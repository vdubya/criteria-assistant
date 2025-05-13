import os
import re
import zipfile
import requests
import pandas as pd
from lxml import etree
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl import Workbook

# --- Global Debug Toggle ---
DEBUG_MODE = False

def debug_print(msg):
    if DEBUG_MODE:
        print(f"[DEBUG] {msg}")

# --- Load CSVs ---
def load_csv_with_required_columns(path, required_cols):
    df = pd.read_csv(path, comment="#")
    df.columns = [c.strip() for c in df.columns]
    for col in required_cols:
        if col not in df.columns:
            raise ValueError(f"Missing required column: '{col}' in {path}")
    return df

def load_discipline_map(path):
    return load_csv_with_required_columns(path, ['UFGS', 'Discipline'])

def load_keyword_list(path):
    return load_csv_with_required_columns(path, ['ID', 'Keyword', 'Recommended Replacement', 'Reason', 'Reference'])

# --- Download and Extract ---
def download_zip(url, zip_path, force=False):
    if os.path.exists(zip_path) and not force:
        debug_print(f"Zip exists at {zip_path}")
        return
    debug_print(f"Downloading from {url}...")
    resp = requests.get(url)
    resp.raise_for_status()
    with open(zip_path, 'wb') as f:
        f.write(resp.content)
    debug_print("Download complete.")

def extract_zip(zip_path, extract_dir):
    if os.path.exists(extract_dir):
        import shutil
        shutil.rmtree(extract_dir)
    with zipfile.ZipFile(zip_path, 'r') as zip_ref:
        zip_ref.extractall(extract_dir)
    debug_print("Extraction complete.")

# --- XML Utilities ---
def clean_text(text):
    return re.sub(r'[^\x20-\x7E]+', '', text.strip()) if text else ''

def extract_text(el):
    return clean_text(''.join(el.itertext())) if el is not None else ''

def get_line_number(el):
    return el.sourceline if hasattr(el, 'sourceline') else -1

def parse_children(el, depth, num_str, ufgs_id):
    rows = []
    for ch in el:
        if ch.tag in ['ENG', 'MET']:
            rows.append((depth+1, ch.tag, num_str, extract_text(ch), get_line_number(ch), None, None, ufgs_id))
        elif len(ch):
            rows += parse_children(ch, depth+1, num_str, ufgs_id)
    return rows

def parse_spt(el, depth, numbering, ufgs_id, submittal=None):
    rows = []
    num_str = '.'.join(map(str, numbering))
    rows.append((depth, 'SPT', num_str, extract_text(el.find('TTL')), get_line_number(el.find('TTL')), None, None, ufgs_id))
    for ch in el:
        tag = ch.tag
        if tag == 'SCP':
            rows.append((depth+1, 'SCP', num_str, extract_text(ch), get_line_number(ch), None, None, ufgs_id))
        elif tag == 'NTE':
            for npr in ch.findall('./NPR'):
                rows.append((depth+1, 'NPR', num_str, extract_text(npr), get_line_number(npr), None, None, ufgs_id))
        elif tag == 'TXT':
            rows.append((depth+1, 'TXT', num_str, extract_text(ch), get_line_number(ch), None, None, ufgs_id))
            rows += parse_children(ch, depth+1, num_str, ufgs_id)
        elif tag == 'LST':
            sub = ch.find('SUB')
            val = extract_text(sub)
            match = re.match(r"(SD-[0-9]{2})", val)
            submittal = match.group(1) if match else None
            rows.append((depth+1, 'LST', num_str, val, get_line_number(ch), submittal, None, ufgs_id))
        elif tag == 'ITM':
            subs = ch.findall('.//SUB')
            tag = 'SUB' if submittal else 'ITM'
            name = extract_text(subs[0]) if subs else extract_text(ch)
            clss = extract_text(subs[1]) if len(subs) > 1 else ''
            itm_depth = depth + 2 if submittal else depth + 1
            rows.append((itm_depth, tag, num_str, name, get_line_number(ch), submittal, clss, ufgs_id))
        elif tag == 'SPT':
            new_num = numbering + [len(el.findall('./SPT')) + 1]
            rows += parse_spt(ch, depth+1, new_num, ufgs_id, submittal)
    return rows

def parse_sec_file(path, ufgs_id):
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        raw = re.sub(r'[^\x09\x0A\x0D\x20-\x7E]', '', f.read())
    root = etree.fromstring(raw.encode('utf-8'))
    rows = []
    for tag in ['SCN', 'STL', 'DTE', 'PRA']:
        for el in root.findall(f".//{tag}"):
            rows.append((0, tag, 'Heading', extract_text(el), get_line_number(el), None, None, ufgs_id))
    for p_idx, prt in enumerate(root.findall('.//PRT'), 1):
        ttl = extract_text(prt.find('TTL'))
        rows.append((1, 'PRT', str(p_idx), ttl, get_line_number(prt.find('TTL')), None, None, ufgs_id))
        for s_idx, spt in enumerate(prt.findall('./SPT'), 1):
            rows += parse_spt(spt, 2, [p_idx, s_idx], ufgs_id)
    return rows

# --- Keyword Warning ---
def keyword_warning(val, df_keywords):
    alerts = []
    for _, row in df_keywords.iterrows():
        if re.search(rf"\b{re.escape(row['Keyword'])}\b", str(val), re.IGNORECASE):
            replacement = f" — use '{row['Recommended Replacement']}'" if row['Recommended Replacement'] else ""
            alerts.append(f"Avoid '{row['Keyword']}'{replacement} [{row['ID']}]")
    return "; ".join(alerts) if alerts else ""

# --- Main Parser ---
def parse_all_sec_files(zip_url, zip_path, extract_dir, discipline_csv_path, keyword_csv_path,
                        output_excel_path=None, include_keyword_warnings=True,
                        debug_mode=False, force_download=False):

    global DEBUG_MODE
    DEBUG_MODE = debug_mode

    # Step 1: Download + Extract
    download_zip(zip_url, zip_path, force=force_download)
    extract_zip(zip_path, extract_dir)

    # Step 2: Load discipline + keyword CSVs
    df_discipline = load_discipline_map(discipline_csv_path)
    df_keywords = load_keyword_list(keyword_csv_path) if include_keyword_warnings else pd.DataFrame()

    # Step 3: Parse all SEC files
    combined, missing = [], []
    ufgs_sections = {f"{row['UFGS']}.SEC": row['UFGS'] for _, row in df_discipline.iterrows()}

    for file, ufgs in ufgs_sections.items():
        path = os.path.join(extract_dir, file)
        if os.path.exists(path):
            combined += parse_sec_file(path, ufgs)
        else:
            debug_print(f"Missing file: {file}")
            missing.append(file)

    # Step 4: Build DataFrame
    df = pd.DataFrame(combined, columns=[
        'NEST_DEPTH', 'SEC TAG', 'Number', 'Value', 'LINE_NUM',
        'SUBMITTAL_CODE', 'Submittal Classification Code', 'UFGS'
    ])
    df.insert(0, 'ID', range(1, len(df) + 1))
    df['SEC TAG'] = df['SEC TAG'].apply(lambda x: f"<{x}>")
    df['TAG LABEL'] = df.apply(lambda r: (
        'Heading' if r['NEST_DEPTH'] == 0 else
        'PART' if r['SEC TAG'] == '<PRT>' else
        'ARTICLE' if r['SEC TAG'] == '<SPT>' and r['Number'].count('.') == 1 else
        'Paragraph' if r['SEC TAG'] == '<SPT>' and r['Number'].count('.') == 2 else
        'Subparagraph' if r['SEC TAG'] == '<SPT>' else
        'English Measurement Units' if r['SEC TAG'] == '<ENG>' else
        'Metric Measurement Units' if r['SEC TAG'] == '<MET>' else
        'Preparing Activity' if r['SEC TAG'] == '<PRA>' else
        'Section Scope' if r['SEC TAG'] == '<SCP>' else ''
    ), axis=1)

    if include_keyword_warnings:
        df["Keyword Warning"] = df["Value"].apply(lambda val: keyword_warning(val, df_keywords))

    # Step 5: Export Excel
    if output_excel_path:
        with pd.ExcelWriter(output_excel_path, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name="ParsedData")
            df_discipline.to_excel(writer, index=False, sheet_name="DisciplineMap")
            if include_keyword_warnings:
                df_keywords.to_excel(writer, index=False, sheet_name="KeywordGuidance")

            for sheet in writer.book.worksheets:
                ws = sheet
                max_col = chr(64 + len(ws[1]))
                tbl_range = f"A1:{max_col}{len(ws['A'])}"
                tbl = Table(displayName=f"{ws.title}Table", ref=tbl_range)
                style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=True)
                tbl.tableStyleInfo = style
                ws.add_table(tbl)
                ws.freeze_panes = "A2"

    return df, df_discipline, {"missing_files": missing}
